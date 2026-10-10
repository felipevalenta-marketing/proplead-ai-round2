from __future__ import annotations

import math
import os
import re
import shutil
import tempfile
from dataclasses import dataclass
from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageColor, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent
CONTENT_MD = ROOT / "final_presentation_content.md"
RUNBOOK_MD = ROOT / "demo_runbook.md"
LANGSMITH_MD = ROOT.parent / "evaluation" / "langsmith.md"
ROI_MD = ROOT.parent / "roi_risk_assessment.md"
OUTPUT_PPTX = ROOT / "PropLead_AI_Final_Presentation.pptx"
BUILD_DIR = ROOT / "_build_final_presentation"
PREVIEW_DIR = BUILD_DIR / "previews"
CONTACT_SHEET = BUILD_DIR / "contact_sheet.png"

W, H = 1920, 1080
NAVY = "#0F1B2D"
NAVY_2 = "#17324D"
SLATE = "#213B57"
CARD = "#F5F0E8"
CARD_2 = "#EDE6DA"
TEXT = "#102033"
TEXT_SOFT = "#35506B"
WHITE = "#F9F6F1"
TEAL = "#2FA8A3"
TEAL_2 = "#4BB3AA"
GREEN = "#5FA86A"
AMBER = "#CC8C3A"
RED = "#B95A53"
GOLD = "#D7B36A"
MUTED = "#7A8EA2"


def font_path(candidates: list[str]) -> str:
    for candidate in candidates:
        if Path(candidate).exists():
            return candidate
    raise FileNotFoundError(f"No font found in {candidates}")

FONT_REG = font_path([
    r"C:\Windows\Fonts\segoeui.ttf",
    r"C:\Windows\Fonts\arial.ttf",
])
FONT_BOLD = font_path([
    r"C:\Windows\Fonts\segoeuib.ttf",
    r"C:\Windows\Fonts\arialbd.ttf",
])
FONT_SEMIBOLD = FONT_BOLD
FONT_ITALIC = font_path([
    r"C:\Windows\Fonts\segoeuii.ttf",
    r"C:\Windows\Fonts\ariali.ttf",
    FONT_REG,
])


def f(size: int, bold: bool = False, italic: bool = False):
    if bold:
        return ImageFont.truetype(FONT_BOLD, size)
    if italic:
        return ImageFont.truetype(FONT_ITALIC, size)
    return ImageFont.truetype(FONT_REG, size)


def hexrgb(color: str):
    return ImageColor.getrgb(color)


def make_canvas(bg=NAVY):
    return Image.new("RGBA", (W, H), hexrgb(bg) + (255,))


def draw_round_rect(draw: ImageDraw.ImageDraw, box, fill, outline=None, width=1, radius=28):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def text_size(draw, text, font):
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=6)
    return bbox[2] - bbox[0], bbox[3] - bbox[1]


def wrap_text(draw, text, font, max_width):
    words = text.split()
    lines = []
    current = ""
    for word in words:
        trial = word if not current else current + " " + word
        if draw.textbbox((0, 0), trial, font=font)[2] <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines or [""]


def draw_wrapped(draw, text, xy, font, fill, max_width, spacing=8, line_gap=10):
    x, y = xy
    lines = []
    for para in text.split("\n"):
        if not para.strip():
            lines.append("")
            continue
        lines.extend(wrap_text(draw, para, font, max_width))
    cur_y = y
    for line in lines:
        draw.text((x, cur_y), line, font=font, fill=fill)
        cur_y += font.size + line_gap
    return cur_y


def add_footer(draw, slide_no: int):
    draw.line((80, 1000, 1840, 1000), fill=hexrgb("#284764"), width=2)
    draw.text((80, 1015), "PropLead AI • Ironhack Round 2", font=f(22), fill=hexrgb(MUTED))
    slide_tag = f"{slide_no:02d}"
    box = (1800, 1000, 1840, 1040)
    draw_round_rect(draw, box, fill=hexrgb(SLATE), outline=hexrgb(TEAL), width=2, radius=12)
    bbox = draw.textbbox((0, 0), slide_tag, font=f(18, bold=True))
    tx = box[0] + (box[2] - box[0] - (bbox[2] - bbox[0])) / 2
    ty = box[1] + (box[3] - box[1] - (bbox[3] - bbox[1])) / 2 - 3
    draw.text((tx, ty), slide_tag, font=f(18, bold=True), fill=hexrgb(WHITE))


def title_block(draw, title, subtitle=None):
    draw.text((90, 70), title, font=f(50, bold=True), fill=hexrgb(WHITE))
    draw.line((90, 145, 520, 145), fill=hexrgb(TEAL), width=5)
    if subtitle:
        draw.text((90, 160), subtitle, font=f(24), fill=hexrgb("#D6E1EA"))


def badge(draw, x, y, text, fill=SLATE, outline=TEAL, width=2, text_fill=WHITE, pad_x=24, pad_y=12, radius=22, font_size=20):
    font = f(font_size, bold=True)
    bbox = draw.textbbox((0, 0), text, font=font)
    w = bbox[2] - bbox[0] + pad_x * 2
    h = bbox[3] - bbox[1] + pad_y * 2
    draw_round_rect(draw, (x, y, x + w, y + h), fill=hexrgb(fill), outline=hexrgb(outline), width=width, radius=radius)
    draw.text((x + pad_x, y + pad_y - 2), text, font=font, fill=hexrgb(text_fill))
    return w, h


def card(draw, box, title, body_lines=None, fill=CARD, outline=TEAL, title_fill=TEXT, body_fill=TEXT_SOFT, title_size=28, body_size=22, radius=28, body_top_offset=60):
    x1, y1, x2, y2 = box
    draw_round_rect(draw, box, fill=hexrgb(fill), outline=hexrgb(outline), width=3, radius=radius)
    draw.text((x1 + 28, y1 + 22), title, font=f(title_size, bold=True), fill=hexrgb(title_fill))
    if body_lines:
        y = y1 + body_top_offset
        for line in body_lines:
            wrapped = wrap_text(draw, line, f(body_size), x2 - x1 - 56)
            for wrapped_line in wrapped:
                draw.text((x1 + 28, y), wrapped_line, font=f(body_size), fill=hexrgb(body_fill))
                y += body_size + 8
            y += 8


def metric_card(draw, box, label, value, accent=TEAL, value_size=34, label_size=18, value_fill=TEXT, label_fill=TEXT_SOFT):
    x1, y1, x2, y2 = box
    draw_round_rect(draw, box, fill=hexrgb(CARD), outline=hexrgb(accent), width=3, radius=24)
    draw.text((x1 + 22, y1 + 18), label, font=f(label_size, bold=True), fill=hexrgb(label_fill))
    draw.text((x1 + 22, y1 + 50), value, font=f(value_size, bold=True), fill=hexrgb(value_fill))


def placeholder(draw, box, title, subtitle=None):
    x1, y1, x2, y2 = box
    draw_round_rect(draw, box, fill=hexrgb("#172A42"), outline=hexrgb(TEAL), width=4, radius=26)
    # dashed border accents
    for i in range(x1 + 20, x2 - 20, 28):
        draw.line((i, y1 + 12, min(i + 14, x2 - 20), y1 + 12), fill=hexrgb(TEAL_2), width=2)
        draw.line((i, y2 - 12, min(i + 14, x2 - 20), y2 - 12), fill=hexrgb(TEAL_2), width=2)
    for j in range(y1 + 20, y2 - 20, 28):
        draw.line((x1 + 12, j, x1 + 12, min(j + 14, y2 - 20)), fill=hexrgb(TEAL_2), width=2)
        draw.line((x2 - 12, j, x2 - 12, min(j + 14, y2 - 20)), fill=hexrgb(TEAL_2), width=2)
    tw = draw.textbbox((0, 0), title, font=f(30, bold=True))
    draw.text((x1 + (x2 - x1 - (tw[2] - tw[0])) / 2, y1 + 80), title, font=f(30, bold=True), fill=hexrgb(WHITE))
    if subtitle:
        lines = wrap_text(draw, subtitle, f(22), x2 - x1 - 120)
        cur_y = y1 + 140
        for line in lines:
            bbox = draw.textbbox((0, 0), line, font=f(22))
            draw.text((x1 + (x2 - x1 - (bbox[2] - bbox[0])) / 2, cur_y), line, font=f(22), fill=hexrgb("#D6E1EA"))
            cur_y += 34


def arrow(draw, x1, y1, x2, y2, color=TEAL, width=6):
    draw.line((x1, y1, x2, y2), fill=hexrgb(color), width=width)
    angle = math.atan2(y2 - y1, x2 - x1)
    head_len = 18
    a1 = angle + math.pi * 0.85
    a2 = angle - math.pi * 0.85
    p1 = (x2 + head_len * math.cos(a1), y2 + head_len * math.sin(a1))
    p2 = (x2 + head_len * math.cos(a2), y2 + head_len * math.sin(a2))
    draw.polygon([(x2, y2), p1, p2], fill=hexrgb(color))


def process_box(draw, box, title, subtitle, accent=TEAL, fill=CARD):
    x1, y1, x2, y2 = box
    draw_round_rect(draw, box, fill=hexrgb(fill), outline=hexrgb(accent), width=3, radius=24)
    draw.text((x1 + 20, y1 + 18), title, font=f(24, bold=True), fill=hexrgb(TEXT))
    lines = wrap_text(draw, subtitle, f(18), x2 - x1 - 40)
    cur_y = y1 + 62
    for line in lines:
        draw.text((x1 + 20, cur_y), line, font=f(18), fill=hexrgb(TEXT_SOFT))
        cur_y += 28


def metric_strip(draw, start_x, y, items, card_w=250, accent=TEAL):
    x = start_x
    for label, value in items:
        metric_card(draw, (x, y, x + card_w, y + 112), label, value, accent=accent, value_size=30, label_size=16)
        x += card_w + 18


def load_sections(md_path: Path):
    text = md_path.read_text(encoding="utf-8")
    pattern = re.compile(r"^##\s+(\d+)\)\s+(.+)$", re.M)
    matches = list(pattern.finditer(text))
    sections = []
    for i, match in enumerate(matches):
        start = match.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        body = text[start:end]
        sec_num = int(match.group(1))
        section_heading = match.group(2).strip()

        def extract(marker: str, next_marker: str | None = None):
            try:
                a = body.index(marker) + len(marker)
            except ValueError:
                return ""
            if next_marker:
                try:
                    b = body.index(next_marker, a)
                except ValueError:
                    b = len(body)
            else:
                b = len(body)
            return body[a:b].strip()

        title = re.search(r"\*\*Slide title:\*\*\s*(.+)", body)
        title_text = title.group(1).strip() if title else section_heading

        on_slide = extract("**On-slide content:**", "**Recommended visual or screenshot:**")
        visual = extract("**Recommended visual or screenshot:**", "**Speaker notes:**")
        notes = extract("**Speaker notes:**", "**What to show on screen:**")
        show = extract("**What to show on screen:**", "**Approx. speaking time:**")
        time = extract("**Approx. speaking time:**")

        bullets = []
        for line in on_slide.splitlines():
            line = line.strip()
            if line.startswith("-"):
                bullets.append(line[1:].strip())
        sections.append({
            "num": sec_num,
            "heading": section_heading,
            "title": title_text,
            "bullets": bullets,
            "visual": visual,
            "notes": notes,
            "show": show,
            "time": time,
        })
    return sections


def parse_langsmith_metrics():
    text = LANGSMITH_MD.read_text(encoding="utf-8")
    def section(name):
        m = re.search(rf"^##\s+{re.escape(name)}\s*$", text, re.M)
        if not m:
            return ""
        start = m.end()
        next_match = re.search(r"^##\s+", text[start:], re.M)
        end = start + next_match.start() if next_match else len(text)
        return text[start:end]

    def values(sec):
        result = {}
        for line in sec.splitlines():
            m = re.match(r"-\s+([^:]+):\s+(.+)$", line.strip())
            if m:
                result[m.group(1).strip()] = m.group(2).strip()
        return result

    return values(section("Structured v1")), values(section("Final structured v2"))


def parse_roi_assumptions():
    text = ROI_MD.read_text(encoding="utf-8")
    vals = {}
    for line in text.splitlines():
        m = re.match(r"\|\s*([^|]+?)\s*\|\s*([^|]+?)\s*\|\s*([^|]+?)\s*\|$", line)
        if m and m.group(1).strip() not in {"Assumption", "Scenario", "Risk", "Record"}:
            vals[m.group(1).strip()] = m.group(2).strip()
    return vals


@dataclass
class SlideAsset:
    image: Path
    notes: str


def build_slide_images(build_dir: Path) -> list[SlideAsset]:
    build_dir.mkdir(parents=True, exist_ok=True)
    preview_dir = build_dir / "previews"
    preview_dir.mkdir(parents=True, exist_ok=True)

    sections = load_sections(CONTENT_MD)
    v1_metrics, v2_metrics = parse_langsmith_metrics()
    roi_vals = parse_roi_assumptions()

    assets = []
    for sec in sections:
        img = make_canvas()
        draw = ImageDraw.Draw(img)
        title_block(draw, sec["title"], "PropLead AI • Ironhack Round 2")
        badge(draw, 1490, 62, "Deployed MVP + POC + Evaluation", fill=SLATE, outline=TEAL)
        if sec["num"] == 1:
            draw.text((90, 245), "Multilingual lead qualification for a Mallorca microagency", font=f(30), fill=hexrgb("#D6E1EA"))
            metric_card(draw, (95, 340, 560, 500), "Browser MVP", "Deterministic\nreview-gated", accent=TEAL)
            metric_card(draw, (95, 530, 560, 690), "n8n POC", "Imported and\nexecuted", accent=GREEN)
            metric_card(draw, (95, 720, 560, 880), "LangSmith", "Hybrid v2\nevaluated", accent=AMBER)
            placeholder(draw, (680, 220, 1810, 890), "Mallorca lead review dashboard", "Use this as the opening visual frame for the capstone")
            badge(draw, 700, 930, "Mandatory human approval", fill=SLATE, outline=GREEN)
            badge(draw, 1080, 930, "Synthetic data documented", fill=SLATE, outline=TEAL_2)
            badge(draw, 1440, 930, "No automatic sending", fill=SLATE, outline=AMBER)
        elif sec["num"] == 2:
            card(draw, (90, 220, 980, 910), "What makes the problem hard", [
                "• Enquiries arrive in English, German and Spanish.",
                "• Agents still need to translate, structure and qualify manually.",
                "• Missing budget, location or property details slow the reply.",
                "• Matching has to stay catalogue-bound and safe.",
            ], fill=CARD, outline=TEAL, title_fill=TEXT, body_fill=TEXT_SOFT, title_size=30, body_size=24, body_top_offset=80)
            metric_card(draw, (1050, 240, 1810, 390), "Languages handled", "3", accent=TEAL)
            metric_card(draw, (1050, 415, 1810, 565), "Agency profile", "Small Mallorca team", accent=GREEN)
            metric_card(draw, (1050, 590, 1810, 740), "Workflow need", "One review screen", accent=AMBER)
            metric_card(draw, (1050, 765, 1810, 915), "Decision goal", "Safer first reply", accent=TEAL_2)
        elif sec["num"] == 3:
            card(draw, (80, 210, 865, 920), "Round 1 feedback", [
                "• Channel definition was still too loose.",
                "• Language consistency needed stronger testing.",
                "• Matching risk and unsafe inference were concerns.",
                "• Synthetic evidence had to be documented clearly.",
            ], fill=CARD, outline=AMBER, title_fill=TEXT, body_fill=TEXT_SOFT, title_size=30, body_size=24, body_top_offset=80)
            card(draw, (1055, 210, 1840, 920), "Round 2 improvements", [
                "• Keep the same use case, but narrow and strengthen it.",
                "• Standardise intake and qualification before matching.",
                "• Keep deterministic catalogue-bound matching.",
                "• Keep mandatory human review for every reply.",
                "• Document the synthetic dataset and final evaluation.",
            ], fill=CARD, outline=GREEN, title_fill=TEXT, body_fill=TEXT_SOFT, title_size=30, body_size=24, body_top_offset=80)
            arrow(draw, 885, 560, 1040, 560, color=TEAL, width=10)
            badge(draw, 790, 905, "KEEP", fill=SLATE, outline=TEAL, font_size=28)
        elif sec["num"] == 4:
            process_box(draw, (80, 300, 610, 820), "Public browser MVP", "Deterministic, review-gated browser interface. Uses synthetic property data and does not call OpenAI.", accent=TEAL)
            process_box(draw, (655, 300, 1285, 820), "n8n operational POC", "Imported seven-node workflow in the target n8n environment. Demonstrates the operational path with simulated intake.", accent=GREEN)
            process_box(draw, (1330, 300, 1840, 820), "LangSmith hybrid extractor", "OpenAI hybrid extraction with deterministic policy. Evaluated separately from the public MVP.", accent=AMBER)
            badge(draw, 150, 860, "Separate layers: browser MVP • n8n POC • LangSmith evaluation", fill=SLATE, outline=TEAL, font_size=22)
        elif sec["num"] == 5:
            # workflow flow
            boxes = [
                (70, 400, 300, 540, "Simulated intake", "web form / email / WhatsApp / portal / social / manual"),
                (335, 400, 565, 540, "Language detection", "EN / DE / ES"),
                (600, 400, 830, 540, "Structured extraction", "Budget • location • type • bedrooms • preferences"),
                (865, 400, 1095, 540, "Qualification gate", "Needs information or ready for matching"),
                (1130, 400, 1360, 540, "Catalogue-bound matching", "Availability and hard filters only"),
                (1395, 400, 1625, 540, "Agent review", "Draft response and approval"),
                (1660, 400, 1850, 540, "Send?", "Never automatic"),
            ]
            for i, (x1, y1, x2, y2, title, subtitle) in enumerate(boxes):
                process_box(draw, (x1, y1, x2, y2), title, subtitle, accent=[TEAL, GREEN, TEAL_2, AMBER, TEAL, GREEN, RED][i])
                if i < len(boxes) - 1:
                    arrow(draw, x2 + 10, 470, boxes[i + 1][0] - 10, 470, color=TEAL, width=6)
            badge(draw, 670, 665, "No automatic customer messaging", fill=SLATE, outline=AMBER, font_size=24)
            card(draw, (140, 720, 1780, 900), "Controlled properties", [
                "Catalogue-bound matching uses a fixed property list, current availability and explicit hard constraints only.",
            ], fill=CARD, outline=TEAL, title_fill=TEXT, body_fill=TEXT_SOFT, title_size=28, body_size=24, body_top_offset=64)
            # Maybe highlight simulated channels
            badge(draw, 140, 930, "Supported intake sources are simulated/standardised channels, not production integrations.", fill=SLATE, outline=TEAL, font_size=18)
        elif sec["num"] == 6:
            placeholder(draw, (85, 230, 1140, 900), "PLACEHOLDER — deployed MVP screenshot", "Insert a screenshot of the live browser MVP at https://proplead-ai-round2-mvp.vercel.app/")
            card(draw, (1200, 240, 1835, 420), "What to point out", [
                "• Deterministic browser MVP.",
                "• Synthetic property data only.",
                "• No customer message is sent automatically.",
            ], fill=CARD, outline=TEAL, title_fill=TEXT, body_fill=TEXT_SOFT, title_size=28, body_size=22, body_top_offset=72)
            metric_card(draw, (1200, 465, 1835, 575), "Human-review gate", "Required", accent=GREEN)
            metric_card(draw, (1200, 600, 1835, 710), "Match type", "Catalogue-bound", accent=TEAL)
            metric_card(draw, (1200, 735, 1835, 845), "Data", "Documented synthetic", accent=AMBER)
            card(draw, (1200, 870, 1835, 960), "Demo input", ["Use the Spanish Port de Sóller enquiry from the runbook."], fill=CARD, outline=GREEN, title_fill=TEXT, body_fill=TEXT_SOFT, title_size=22, body_size=18, body_top_offset=48)
        elif sec["num"] == 7:
            placeholder(draw, (85, 230, 1000, 900), "PLACEHOLDER — n8n successful execution screenshot", "Insert the run evidence from the imported seven-node workflow")
            # node row diagram on the right
            node_x = 1070
            node_y = 330
            node_w = 170
            node_h = 86
            labels = ["Manual\ntrigger", "Simulated\ninput", "Normalise", "Extract +\nqualify", "Match", "Draft", "Review"]
            for idx, label in enumerate(labels):
                box = (node_x + idx * 106, node_y + (idx % 2) * 90, node_x + idx * 106 + node_w, node_y + (idx % 2) * 90 + node_h)
                process_box(draw, box, label, "", accent=[TEAL, GREEN, TEAL_2, AMBER, TEAL, GREEN, TEAL_2][idx], fill=CARD)
                if idx < len(labels) - 1:
                    arrow(draw, box[2] + 5, box[1] + 43, node_x + (idx + 1) * 106 - 8, node_y + ((idx + 1) % 2) * 90 + 43, color=TEAL, width=5)
            card(draw, (1070, 620, 1835, 845), "What the screenshot should prove", [
                "• Imported successfully in the target n8n environment.",
                "• Seven-node workflow executed successfully.",
                "• Separate from the browser MVP.",
                "• Still based on simulated/synthetic input.",
            ], fill=CARD, outline=GREEN, title_fill=TEXT, body_fill=TEXT_SOFT, title_size=26, body_size=21, body_top_offset=68)
        elif sec["num"] == 8:
            card(draw, (85, 210, 1835, 890), "LangSmith comparison", None, fill=CARD, outline=TEAL, title_fill=TEXT, title_size=30)
            draw.text((120, 280), "Baseline v1 — structured_extractor_v1-fd0b3bae", font=f(24, bold=True), fill=hexrgb(TEXT))
            draw.text((1060, 280), "Final hybrid v2 — structured_extractor_v2-d2454c34", font=f(24, bold=True), fill=hexrgb(TEXT))
            # Baseline column
            metric_card(draw, (120, 330, 820, 416), "Language accuracy", v1_metrics.get("Language accuracy", "100%"), accent=TEAL, value_size=30)
            metric_card(draw, (120, 430, 820, 516), "Human-review gate", v1_metrics.get("Human-review gate", "100%"), accent=GREEN, value_size=30)
            metric_card(draw, (120, 530, 820, 616), "No-critical-fabrication", v1_metrics.get("No-critical-fabrication", "94.4444%"), accent=AMBER, value_size=28)
            metric_card(draw, (120, 630, 820, 716), "Explicit-field accuracy", v1_metrics.get("Explicit-field accuracy", "92.4%"), accent=AMBER, value_size=28)
            badge(draw, 120, 760, f"Escalation accuracy: {v1_metrics.get('Escalation accuracy', '61.1%')}", fill=SLATE, outline=RED, font_size=20)
            # v2 column
            metric_card(draw, (1060, 330, 1760, 416), "Matching correctness", v2_metrics.get("Matching correctness", "100%"), accent=GREEN, value_size=30)
            metric_card(draw, (1060, 430, 1760, 516), "Escalation correctness", v2_metrics.get("Escalation correctness", "100%"), accent=GREEN, value_size=30)
            metric_card(draw, (1060, 530, 1760, 616), "Human-review gate", v2_metrics.get("Human-review gate", "100%"), accent=GREEN, value_size=30)
            metric_card(draw, (1060, 630, 1760, 716), "Language correctness", v2_metrics.get("Language correctness", "100%"), accent=GREEN, value_size=30)
            badge(draw, 1060, 760, f"Explicit / no-critical: {v2_metrics.get('Explicit-field accuracy', '93.0556%')} / {v2_metrics.get('No-critical-fabrication', '94.4444%')}", fill=SLATE, outline=TEAL, font_size=18)
            badge(draw, 85, 915, "Average latency: 1.5378 seconds", fill=SLATE, outline=TEAL_2, font_size=22)
            badge(draw, 515, 915, "Total tokens: 7,753", fill=SLATE, outline=TEAL_2, font_size=22)
            badge(draw, 810, 915, "Total cost: USD 0.01120725", fill=SLATE, outline=TEAL_2, font_size=22)
            badge(draw, 1320, 915, "v1 revealed the escalation weakness; v2 fixed the matching path", fill=SLATE, outline=AMBER, font_size=18)
        elif sec["num"] == 9:
            card(draw, (80, 220, 640, 900), "ROI assumptions", [
                f"• Agency size: 1–9 employees.",
                f"• Buyer enquiries / month: {roi_vals.get('Buyer enquiries per month', '60')}.",
                f"• Loaded agent cost: {roi_vals.get('Loaded agent cost', '€25/hour')}.",
                f"• Implementation cost: {roi_vals.get('Upfront implementation', '€4,200')}.",
                f"• Operating cost: {roi_vals.get('Monthly operating cost', '€240')}.",
            ], fill=CARD, outline=TEAL, title_fill=TEXT, body_fill=TEXT_SOFT, title_size=28, body_size=23, body_top_offset=84)
            card(draw, (690, 220, 1260, 900), "Main risks", [
                "• Operational: wrong or unavailable property shown.",
                "• Technical: language drift or extraction failure.",
                "• Reputational: unsafe or overconfident reply.",
                "• Compliance: privacy, logging and review gaps.",
                "• Mitigation: human approval, hard filters and review logs.",
            ], fill=CARD, outline=AMBER, title_fill=TEXT, body_fill=TEXT_SOFT, title_size=28, body_size=23, body_top_offset=84)
            card(draw, (1310, 220, 1840, 900), "Responsible AI controls", [
                "• Mandatory human approval.",
                "• Catalogue-bound matching only.",
                "• Synthetic data documented.",
                "• Data minimisation and retention planning.",
                "• Assessment, not legal-compliance claim.",
            ], fill=CARD, outline=GREEN, title_fill=TEXT, body_fill=TEXT_SOFT, title_size=28, body_size=23, body_top_offset=84)
            badge(draw, 90, 930, "ROI values are assumptions/projections, not measured results.", fill=SLATE, outline=AMBER, font_size=20)
        elif sec["num"] == 10:
            # timeline
            steps = [
                (100, 390, 430, 620, "Week 1", "Readiness", "Freeze the schema, controls and pilot rules."),
                (490, 390, 820, 620, "Week 2", "Controlled pilot", "One agency, 2–3 agents, limited channels."),
                (880, 390, 1210, 620, "Week 3", "Weekly review", "Inspect corrections, rejections and rollback triggers."),
                (1270, 390, 1600, 620, "Week 4", "Decision", "Continue, narrow or pause based on the evidence."),
            ]
            for idx, (x1, y1, x2, y2, week, title, body) in enumerate(steps):
                card(draw, (x1, y1, x2, y2), f"{week} — {title}", [f"{body}"], fill=CARD, outline=[TEAL, GREEN, AMBER, TEAL_2][idx], title_fill=TEXT, body_fill=TEXT_SOFT, title_size=24, body_size=22, body_top_offset=72)
                if idx < len(steps) - 1:
                    arrow(draw, x2 + 8, 505, steps[idx + 1][0] - 10, 505, color=TEAL, width=6)
            card(draw, (170, 710, 1750, 930), "Pilot conclusion", [
                "No automatic customer messaging during the pilot. Scale only if safety, usability and ROI support it.",
            ], fill=CARD, outline=TEAL, title_fill=TEXT, body_fill=TEXT_SOFT, title_size=28, body_size=24, body_top_offset=74)
            badge(draw, 170, 950, "Recommended next step: final presentation + safe pilot approval", fill=SLATE, outline=GREEN, font_size=22)
        else:
            draw.text((100, 300), "Unhandled slide", font=f(36), fill=hexrgb(WHITE))

        add_footer(draw, sec["num"])
        img_path = preview_dir / f"slide_{sec['num']:02d}.png"
        img.convert("RGB").save(img_path, quality=95)
        notes = f"{sec['notes']}\n\nWhat to show on screen:\n{sec['show']}\n\nApprox. speaking time: {sec['time']}\n\nRecommended visual: {sec['visual']}"
        assets.append(SlideAsset(img_path, notes))

    # contact sheet
    thumbs = []
    for asset in assets:
        im = Image.open(asset.image).convert("RGB")
        thumb = ImageOps.contain(im, (700, 394), Image.Resampling.LANCZOS)
        framed = Image.new("RGB", (720, 430), (15, 27, 45))
        framed.paste(thumb, ((720 - thumb.width) // 2, 20))
        d = ImageDraw.Draw(framed)
        idx = assets.index(asset) + 1
        d.rectangle((0, 0, 719, 429), outline=(47, 168, 163), width=3)
        d.text((22, 392), f"Slide {idx:02d}", font=f(24, bold=True), fill=(249, 246, 241))
        thumbs.append(framed)
    sheet = Image.new("RGB", (1440, 2150), (15, 27, 45))
    for i, thumb in enumerate(thumbs):
        x = 0 if i % 2 == 0 else 720
        y = (i // 2) * 430
        sheet.paste(thumb, (x, y))
    sheet.save(CONTACT_SHEET)
    return assets


def build_pptx(assets: list[SlideAsset]):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    # remove default slides created by Presentation? none.
    for asset in assets:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(asset.image), 0, 0, width=prs.slide_width, height=prs.slide_height)
        slide.notes_slide.notes_text_frame.text = asset.notes
    # Delete the initial empty slide that python-pptx creates? Presentation() starts with zero slides.
    prs.save(OUTPUT_PPTX)


def main():
    if BUILD_DIR.exists():
        shutil.rmtree(BUILD_DIR)
    assets = build_slide_images(BUILD_DIR)
    build_pptx(assets)
    print(f"Created {OUTPUT_PPTX}")
    print(f"Created {CONTACT_SHEET}")


if __name__ == "__main__":
    main()

